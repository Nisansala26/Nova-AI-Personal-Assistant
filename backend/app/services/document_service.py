import io
import os
import re
from typing import Dict, List, Tuple

import docx
from pypdf import PdfReader

from app.services.ai_service import ask_ai

MAX_FILE_SIZE_BYTES = 10 * 1024 * 1024  # 10 MB
ALLOWED_EXTENSIONS = {".pdf", ".docx", ".txt", ".md", ".json", ".csv"}


class DocumentError(Exception):
    pass


def validate_file(filename: str, file_size: int) -> str:
    """Validate filename, extension, and file size."""
    if file_size > MAX_FILE_SIZE_BYTES:
        raise DocumentError(f"File size exceeds the 10 MB limit ({file_size / (1024*1024):.1f} MB).")

    ext = os.path.splitext(filename)[1].lower()
    if ext not in ALLOWED_EXTENSIONS:
        allowed_list = ", ".join(ALLOWED_EXTENSIONS)
        raise DocumentError(f"Unsupported file type '{ext}'. Allowed extensions: {allowed_list}")

    return ext


def extract_text_from_file(filename: str, file_bytes: bytes) -> Tuple[str, int]:
    """Extract plain text and word count from PDF, DOCX, or text files."""
    ext = validate_file(filename, len(file_bytes))
    text_content = ""

    try:
        if ext == ".pdf":
            reader = PdfReader(io.BytesIO(file_bytes))
            pages_text = []
            for i, page in enumerate(reader.pages):
                extracted = page.extract_text() or ""
                if extracted.strip():
                    pages_text.append(f"--- [Page {i+1}] ---\n{extracted.strip()}")
            text_content = "\n\n".join(pages_text)

        elif ext == ".docx":
            doc = docx.Document(io.BytesIO(file_bytes))
            paras = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
            for table in doc.tables:
                for row in table.rows:
                    row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
                    if row_text:
                        paras.append(row_text)
            text_content = "\n\n".join(paras)

        else:
            # .txt, .md, .json, .csv
            try:
                text_content = file_bytes.decode("utf-8")
            except UnicodeDecodeError:
                text_content = file_bytes.decode("latin-1", errors="replace")

    except Exception as e:
        raise DocumentError(f"Could not extract text from '{filename}': {str(e)}") from e

    cleaned_text = text_content.strip()
    if not cleaned_text:
        raise DocumentError(f"'{filename}' appears to be empty or contains no readable text.")

    words = len(cleaned_text.split())
    return cleaned_text, words


async def generate_document_summary(content: str, filename: str) -> str:
    """Generate an executive AI summary of the document text."""
    # Truncate content to fit in LLM context window while preserving core content
    truncated = content[:12000] if len(content) > 12000 else content

    prompt = (
        f"You are an expert document analysis assistant. Provide a structured, insightful executive summary "
        f"of the following document '{filename}'.\n\n"
        f"Structure your answer with:\n"
        f"1. **Core Overview**: 1-2 sentence executive summary.\n"
        f"2. **Key Findings & Insights**: Bullet points covering the most important concepts.\n"
        f"3. **Action Items / Next Steps** (if applicable).\n\n"
        f"Document Content:\n{truncated}"
    )

    summary = await ask_ai(prompt)
    return summary


async def ask_document_question(content: str, filename: str, question: str) -> Dict[str, Any]:
    """
    Answer a question based on document content, providing direct source references.
    """
    q = question.strip()
    if not q:
        raise DocumentError("Question cannot be empty.")

    # Find relevant paragraphs containing keywords from the question
    keywords = [w.lower() for w in re.findall(r"\w+", q) if len(w) > 3]
    paragraphs = [p.strip() for p in content.split("\n\n") if len(p.strip()) > 30]

    matched_excerpts: List[str] = []
    for p in paragraphs:
        p_lower = p.lower()
        if any(kw in p_lower for kw in keywords):
            matched_excerpts.append(p[:400] + ("..." if len(p) > 400 else ""))
            if len(matched_excerpts) >= 4:
                break

    # If no specific keyword match, take first key paragraphs
    if not matched_excerpts and paragraphs:
        matched_excerpts = [p[:350] for p in paragraphs[:3]]

    # Provide up to 10k chars of context
    context_text = content[:10000]

    prompt = (
        f"You are Nova AI, an expert document analyst. Answer the user's question accurately using ONLY "
        f"the provided document content from '{filename}'.\n"
        f"Always reference specific sections or pages when available.\n"
        f"If the answer cannot be determined from the document, state that clearly.\n\n"
        f"User Question: {q}\n\n"
        f"Document Content:\n{context_text}"
    )

    answer = await ask_ai(prompt)

    return {
        "answer": answer,
        "references": matched_excerpts,
        "filename": filename,
    }
